from docx import Document
from docx.shared import Pt

import docxedit


def make_doc(*paragraphs):
    """Create a Document with the given paragraph texts."""
    doc = Document()
    for text in paragraphs:
        doc.add_paragraph(text)
    return doc


def make_doc_with_table(rows, cols, cell_texts=None):
    """Create a Document with a table. cell_texts is a list of (row, col, text) tuples."""
    doc = Document()
    table = doc.add_table(rows=rows, cols=cols)
    if cell_texts:
        for row, col, text in cell_texts:
            table.cell(row, col).paragraphs[0].add_run(text)
    return doc


class TestReplaceString:
    def test_replaces_in_paragraph(self):
        doc = make_doc("Hello world")
        docxedit.replace_string(doc, "Hello", "Goodbye")
        texts = [p.text for p in doc.paragraphs]
        assert any("Goodbye world" in t for t in texts)

    def test_replaces_multiple_instances(self):
        doc = make_doc("Hello world", "Hello again")
        docxedit.replace_string(doc, "Hello", "Goodbye")
        texts = [p.text for p in doc.paragraphs]
        assert any("Goodbye world" in t for t in texts)
        assert any("Goodbye again" in t for t in texts)

    def test_no_match_leaves_text_unchanged(self):
        doc = make_doc("Hello world")
        docxedit.replace_string(doc, "Foo", "Bar")
        texts = [p.text for p in doc.paragraphs]
        assert any("Hello world" in t for t in texts)
        assert not any("Bar" in t for t in texts)

    def test_replaces_in_table(self):
        doc = make_doc_with_table(2, 2, [(0, 0, "Hello")])
        docxedit.replace_string(doc, "Hello", "Goodbye", include_tables=True)
        assert doc.tables[0].cell(0, 0).text == "Goodbye"

    def test_exclude_tables(self):
        doc = make_doc_with_table(2, 2, [(0, 0, "Hello")])
        docxedit.replace_string(doc, "Hello", "Goodbye", include_tables=False)
        assert doc.tables[0].cell(0, 0).text == "Hello"


class TestReplaceStringUpToParagraph:
    def test_replaces_only_up_to_limit(self):
        doc = make_doc("Hello first", "Hello second", "Hello third")
        para_count = len(doc.paragraphs)
        # Replace in the first 2 paragraphs only
        docxedit.replace_string_up_to_paragraph(doc, "Hello", "Goodbye", paragraph_number=2)
        replaced = [p.text for p in doc.paragraphs if "Goodbye" in p.text]
        not_replaced = [p.text for p in doc.paragraphs if "Hello" in p.text]
        assert len(replaced) >= 1
        assert len(not_replaced) >= 1

    def test_zero_paragraphs_replaces_nothing(self):
        doc = make_doc("Hello first", "Hello second")
        docxedit.replace_string_up_to_paragraph(doc, "Hello", "Goodbye", paragraph_number=0)
        texts = [p.text for p in doc.paragraphs]
        assert not any("Goodbye" in t for t in texts)


class TestRemoveParagraph:
    def test_removes_paragraph(self):
        doc = make_doc("First", "Second", "Third")
        # Find and remove the paragraph with "Second"
        for p in doc.paragraphs:
            if p.text == "Second":
                docxedit.remove_paragraph(p)
                break
        texts = [p.text for p in doc.paragraphs]
        assert "Second" not in texts
        assert "First" in texts
        assert "Third" in texts


class TestRemoveLines:
    def test_removes_matching_line_and_following(self):
        doc = make_doc("Line 1", "Remove this", "Also remove", "Keep this")
        docxedit.remove_lines(doc, first_line="Remove this", number_of_lines=1)
        texts = [p.text for p in doc.paragraphs]
        assert "Remove this" not in texts
        assert "Also remove" not in texts
        assert "Keep this" in texts


class TestAddTextInTable:
    def test_adds_text_to_cell(self):
        doc = make_doc_with_table(2, 2)
        docxedit.add_text_in_table(doc.tables[0], row_num=0, column_num=0, new_string="Hello")
        assert doc.tables[0].cell(0, 0).text == "Hello"

    def test_overwrites_existing_text(self):
        doc = make_doc_with_table(2, 2, [(0, 0, "Old")])
        docxedit.add_text_in_table(doc.tables[0], row_num=0, column_num=0, new_string="New")
        assert doc.tables[0].cell(0, 0).text == "New"


class TestChangeTableFontSize:
    def test_changes_font_size(self):
        doc = make_doc_with_table(1, 1, [(0, 0, "Hello")])
        docxedit.change_table_font_size(doc.tables[0], font_size=14)
        run = doc.tables[0].cell(0, 0).paragraphs[0].runs[0]
        assert run.font.size == Pt(14)
